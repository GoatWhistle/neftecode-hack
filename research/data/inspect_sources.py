"""Read-only inventory and profiling of the supplied hackathon data."""
from pathlib import Path
import json
from collections import Counter
from datetime import datetime
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent

def dump(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str))

def profile_series(dates, values):
    numeric = pd.to_numeric(pd.Series(values), errors='coerce')
    valid_dates = pd.to_datetime(pd.Series(dates), errors='coerce')
    ordered = valid_dates.dropna().sort_values()
    return dict(count=len(values), numeric=int(numeric.notna().sum()),
                nonnumeric=Counter(str(v) for v, n in zip(values, numeric) if pd.isna(n)),
                start=str(ordered.min()), end=str(ordered.max()),
                duplicate_dates=int(valid_dates.dropna().duplicated().sum()),
                nonascending=int((valid_dates.diff().dt.total_seconds() < 0).sum()),
                min=float(numeric.min()) if numeric.notna().any() else None,
                max=float(numeric.max()) if numeric.notna().any() else None,
                median=float(numeric.median()) if numeric.notna().any() else None,
                negatives=int((numeric < 0).sum()), zeros=int((numeric == 0).sum()),
                median_gap_hours=float(ordered.diff().dt.total_seconds().median()/3600),
                max_gap_hours=float(ordered.diff().dt.total_seconds().max()/3600))

def main():
    profiles = {}
    time_grids = {}
    for f in sorted((ROOT/'task/data').glob('*.csv')):
        df = pd.read_csv(f)
        dt = pd.to_datetime(df['date'], errors='coerce')
        time_grids[f.name] = dt
        data = df.drop(columns=[c for c in df if c == 'date' or c.startswith('Unnamed:')])
        numeric = data.apply(pd.to_numeric, errors='coerce')
        per_column = {}
        for c in data:
            s = numeric[c]
            per_column[c] = dict(missing=int(data[c].isna().sum()),
                nonnumeric=int((data[c].notna() & s.isna()).sum()),
                unique=int(s.nunique()), min=float(s.min()), max=float(s.max()),
                median=float(s.median()), p01=float(s.quantile(.01)), p99=float(s.quantile(.99)),
                zeros=int((s == 0).sum()), negatives=int((s < 0).sum()),
                infinite=int(np.isinf(s).sum()))
        profiles[f.name] = dict(rows=len(df), raw_columns=len(df.columns), signals=len(data.columns),
            start=str(dt.min()), end=str(dt.max()), duplicate_dates=int(dt.duplicated().sum()),
            invalid_dates=int(dt.isna().sum()), monotonic=dt.is_monotonic_increasing,
            interval_seconds={str(k): int(v) for k,v in dt.diff().dt.total_seconds().value_counts().items()},
            missing_cells=int(data.isna().sum().sum()), columns=per_column)
        print(f.name, {k:v for k,v in profiles[f.name].items() if k != 'columns'}, flush=True)
    profiles['same_time_grid'] = time_grids['avt_tags.csv'].equals(time_grids['242000_tags.csv'])
    dump('telemetry-profile.json', profiles)

    lab_file = next((ROOT/'task').glob('ЛИМС*.xlsx'))
    wb = openpyxl.load_workbook(lab_file, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.values)
    lab = []
    point = None
    for col in range(0, len(rows[0]), 2):
        point = rows[0][col] or point
        dates, vals, bad = [], [], []
        for ridx, row in enumerate(rows[4:], 5):
            d, v = row[col], row[col+1]
            if d is None and v is None: continue
            if isinstance(d, datetime): dates.append(d); vals.append(v)
            else: bad.append(dict(row=ridx,date=str(d),value=str(v)))
        rec=dict(point=point, indicator=rows[1][col], unit=rows[2][col],
                 cells=f'{get_column_letter(col+1)}:{get_column_letter(col+2)}',
                 declared_count=rows[3][col+1], nondate_rows=bad,
                 **profile_series(dates, vals))
        lab.append(rec)
    dump('laboratory-profile.json',lab)
    print('LIMS series',len(lab),'values',sum(x['count'] for x in lab), flush=True)
    wb.close()

    pak_file = next((ROOT/'task').glob('Выгрузка*.xlsx'))
    wb = openpyxl.load_workbook(pak_file, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    names, units = next(it), next(it)
    columns = [i for i,n in enumerate(names) if n is not None]
    records = {c:dict(dates=[],values=[],bad=[]) for c in columns}
    for ridx,row in enumerate(it,3):
        for c in columns:
            d,v=row[c],row[c+1]
            if d is None and v is None: continue
            if isinstance(d,datetime): records[c]['dates'].append(d);records[c]['values'].append(v)
            else: records[c]['bad'].append(dict(row=ridx,date=str(d),value=str(v)))
    pak=[]
    for c,r in records.items():
        rec=dict(tag=names[c],unit=units[c],cells=f'{get_column_letter(c+1)}:{get_column_letter(c+2)}',
                 nondate_rows=r['bad'], **profile_series(r['dates'],r['values']))
        rec['numeric_gt_10']=int((pd.to_numeric(pd.Series(r['values']),errors='coerce')>10).sum())
        pak.append(rec)
    dump('online-profile.json',pak)
    print('PAK',pak,flush=True)
    wb.close()

    metadata={}
    for f in (ROOT/'task').glob('*.xlsx'):
        w=openpyxl.load_workbook(f,read_only=False,data_only=False)
        metadata[f.name]={s.title:dict(rows=s.max_row,columns=s.max_column,
            merged=[str(x) for x in s.merged_cells.ranges],
            formula_count=sum(c.data_type=='f' for row in s for c in row),
            error_count=sum(c.data_type=='e' for row in s for c in row),
            comments=[dict(cell=c.coordinate,text=c.comment.text) for row in s for c in row if c.comment]) for s in w}
        w.close()
    dump('workbook-metadata.json',metadata)

if __name__ == '__main__':
    main()
